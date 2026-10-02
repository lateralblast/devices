#!/usr/bin/env python3
#
# Name:         devices (Python)
# Version:      0.6.2
# Release:      1
# License:      CC-BA (Creative Commons By Attribution)
#               http://creativecommons.org/licenses/by/4.0/legalcode
# Group:        System
# Source:       N/A
# URL:          N/A
# Distribution: UNIX
# Vendor:       UNIX
# Packager:     Richard Spindler <richard@lateralblast.com.au>
# Description:  Generate front and rear rack elevation diagrams from a CSV file
#               of datacenter hardware, using SVGs extracted from Visio
#               stencils (via devon.py). Python counterpart of devices.ps1 that
#               needs neither Windows nor Visio.

"""Rack diagram generator: CSV in, SVG/PNG/JPG/PDF out.

Stencils come from the visio-stencils repository (zipped .vss files). Each
stencil is split into one SVG per master the first time it is needed, using
devon.py, and the SVGs are cached under svg-cache/.
"""

import argparse
import copy
import csv
import datetime
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
import xml.etree.ElementTree as ET

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
SVG_NS = "http://www.w3.org/2000/svg"
XLINK_NS = "http://www.w3.org/1999/xlink"
ET.register_namespace("", SVG_NS)
ET.register_namespace("xlink", XLINK_NS)


def get_script_vers():
  """Read the version back out of this script's own header comment."""
  with open(os.path.abspath(__file__)) as f:
    for line in f:
      if line.startswith("# Version:"):
        return line.split()[2]
  return "unknown"


# --------------------------------------------------------------------------
# Stencil table (mirrors the $<vendor>_..._stencils_file variables in
# devices.ps1). Paths are relative to the visio-stencils directory.
# --------------------------------------------------------------------------

STENCIL_FILES = {
    "oracle_sparc_server":  "o/oracle/Oracle-Server-SPARC.vss",
    "oracle_intel_server":  "o/oracle/Oracle-Server-x86.vss",
    "oracle_blade_server":  "o/oracle/Oracle-Server-Blade.vss",
    "dell_rack":            "d/dell/Dell-Racks.vss",
    "dell_blade_server":    "d/dell/Dell-PowerEdge-BladeServers.vss",
    "dell_rack_server":     "d/dell/Dell-PowerEdge-RackServers.vss",
    "dell_classic_server":  "d/dell/Dell-PowerEdge-RackServers-Classic.vss",
    "dell_sc_storage":      "d/dell/Dell-Storage-Compellent-SC.vss",
    "dell_ps_storage":      "d/dell/Dell-Storage-EqualLogic-PS.vss",
    "dell_md_storage":      "d/dell/Dell-Storage-PowerVault-Dx-MD-NX.vss",
    "dell_emc_storage":     "d/dell/Dell-EMC.vss",
    "pure_storage_array":   "p/pure/Purestorage.vss",
}

DEFAULT_RACK = "4220 Rack Frame"
BLANK_PLATE = "1U Metal Close Out"

# Dell R/C models in the current rack server stencil; older ones are in the
# Classic stencil. Same list as $dell_current_server_models in devices.ps1.
DELL_CURRENT_SERVER_MODELS = (
    r"^C6520|^C6525|^C6600|^R250|^R350|^R450|^R470|^R550|^R570|^R650|^R660"
    r"|^R670|^R750|^R760|^R770|^R840|^R860|^R940|^R960|^R6525|^R6615|^R6625"
    r"|^R6715|^R6725|^R7525|^R7615|^R7625|^R7715|^R7725"
)

# --------------------------------------------------------------------------
# Layout constants, in inches (same values as the devices.ps1 constants)
# --------------------------------------------------------------------------

RU_SPACE = 0.175         # height of one rack unit
RACK_PITCH = 4.0         # front rack x to rear rack x ($back_rack_x - $front_rack_x)
RU_BASE = 0.30           # rack frame bottom edge to the bottom of RU 1
MARGIN_X = 0.40
MARGIN_Y = 0.15
PAGE_LABEL_HEIGHT = 0.5  # extra room at the top for -pagelabels
UNITS = 100.0            # SVG user units per inch in the composed drawing

LABEL_FONT_PT = 4.8      # devices.ps1 asks for 6pt, but Visio renders it this size in rack.jpg
LABEL_FILL = "#f5a623"
LABEL_TEXT = "#000000"


def csi(pattern, text):
  """PowerShell -match / switch -regex are case insensitive."""
  return re.search(pattern, text or "", re.IGNORECASE) is not None


# --------------------------------------------------------------------------
# Vendor/model dispatch (mirrors the switch -regex block in devices.ps1)
# --------------------------------------------------------------------------

def pick_shape(row):
  """Return (stencil_key, front_master, back_master) for a CSV row.

  Like switch -regex, every matching case runs and a later match overrides
  an earlier one; default only applies when nothing matched.
  """
  vendor = row["Vendor"]
  model = row["Model"]
  arch = row["Architecture"]
  blank = ("dell_rack", BLANK_PLATE, BLANK_PLATE)

  if csi("Dell", vendor):
    front, back = f"{model} Front", f"{model} Rear"
    result = None
    if csi(r"^CX4|^NX4|^ES|^DD", model):
      result = "dell_emc_storage"
    if csi(DELL_CURRENT_SERVER_MODELS, model):
      result = "dell_rack_server"
    # Complement of the current-stencil case; CX4/NX4 have their own case
    if csi(r"^R|^C", model) and not csi(DELL_CURRENT_SERVER_MODELS, model) \
            and not csi(r"^CX4|^NX4", model):
      result = "dell_classic_server"
    if csi(r"^M[0-9]", model):
      result = "dell_blade_server"
    if csi(r"^FS8|^SC", model):
      result = "dell_sc_storage"
    if csi(r"^FS7|^PS", model):
      result = "dell_ps_storage"
    if csi(r"^D|^MD|^NX", model):
      result = "dell_md_storage"
    if result is None:
      return blank
    return (result, front, back)

  if csi("Pure", vendor):
    front, back = f"{model} front", f"{model} back"
    if csi(r"FB|FlashBlade", model):
      front, back = "FlashBlade Front Full", "FlashBlade back"
    if csi(r"FA-|M", model):
      front, back = "FA M70 front", "FA M70 back"
    return ("pure_storage_array", front, back)

  if csi("Oracle|Sun", vendor):
    front, back = f"{model} Front", f"{model} Rear"
    if csi("Blade", model):
      return ("oracle_blade_server", front, back)
    if csi("SPARC", arch):
      return ("oracle_sparc_server", front, back)
    return ("oracle_intel_server", front, back)

  return blank


# --------------------------------------------------------------------------
# Stencil -> SVG extraction and lookup
# --------------------------------------------------------------------------

def shape_filename(master):
  """Same sanitising as devon.safe_shape_filename (without the duplicate suffix)."""
  return (re.sub(r"[^A-Za-z0-9._-]+", "_", master or "").strip("_") or "shape") + ".svg"


# Vendor names as they appear in a CSV, mapped to the directory name in visio-stencils
VENDOR_ALIASES = {"hp": "hpe", "hewlettpackard": "hpe", "hewlettpackardenterprise": "hpe", "sun": "oracle",
                  "sunmicrosystems": "oracle", "dellemc": "dell", "emc": "dell", "lenovoibm": "lenovo"}


def norm(text):
  return re.sub(r"[^a-z0-9]", "", (text or "").lower())


def model_pattern(model):
  """Regex for a model name that ignores spacing/punctuation differences (DL380 Gen10, DL380-Gen10)."""
  chunks = re.findall(r"[A-Za-z0-9]+", model or "")
  return r"[\W_]*".join(re.escape(c) for c in chunks)


def best_master(names, model, words):
  """Best master name for a model whose name mentions one of words (e.g. front, rear/back).
  Preference: '<model> <word>', then the model as a whole word at the start (R650 4D Front but
  not R650xs), then as a prefix, then anywhere; open/door/bezel variants and longer names lose ties."""
  pat = model_pattern(model)
  if not pat:
    return None
  word_re = "|".join(words)
  best = None
  for name in names:
    low = name.lower()
    if not any(w in low for w in words):
      continue
    if re.match(rf"^{pat}[\W_]*(?:{word_re})\w*$", name, re.IGNORECASE):
      tier = 0
    elif re.match(rf"^{pat}(?![A-Za-z0-9])", name, re.IGNORECASE):
      tier = 1
    elif re.match(pat, name, re.IGNORECASE):
      tier = 2
    elif re.search(pat, name, re.IGNORECASE):
      tier = 3
    else:
      continue
    score = (tier, bool(re.search(r"open|door|bezel|empty|blank", low)), len(name))
    if best is None or score < best[0]:
      best = (score, name)
  return best[1] if best else None


class StencilLibrary:
  def __init__(self, stencil_dir, cache_dir, devon, verbose=False, max_scan=30):
    self.max_scan = max_scan
    self.found = {}      # (vendor, model) -> discovered (key, front, back) or None
    self.stencil_dir = stencil_dir
    self.cache_dir = cache_dir
    self.devon = devon
    self.verbose = verbose
    self.dirs = {}       # stencil key -> directory of SVGs
    self.parsed = {}     # svg path -> (root element, width_in, height_in)

  def log(self, msg):
    if self.verbose:
      print(msg, file=sys.stderr)

  def find_devon(self):
    candidates = [
        self.devon,
        os.environ.get("DEVON"),
        os.path.join(SCRIPT_DIR, "..", "devon", "devon.py"),
        shutil.which("devon.py"),
    ]
    for cand in candidates:
      if cand and os.path.isfile(cand):
        return os.path.abspath(cand)
    sys.exit("devon.py not found; use --devon PATH or set DEVON (see https://github.com/lateralblast/devon)")

  def svg_dir(self, key):
    """Directory of per-master SVGs for a stencil, extracting them on first use."""
    if key in self.dirs:
      return self.dirs[key]
    rel = STENCIL_FILES.get(key, key)   # discovered stencils use their relative path as the key
    stencil_file = os.path.join(self.stencil_dir, *rel.split("/"))
    stem = os.path.splitext(os.path.basename(rel))[0]
    out_dir = os.path.join(self.cache_dir, stem)
    done = os.path.join(out_dir, ".done")
    if not os.path.exists(done):
      self.extract(stencil_file, out_dir, stem)
      open(done, "w").close()
    self.dirs[key] = out_dir
    return out_dir

  def unpack(self, stencil_file, work):
    """Return a readable stencil file, extracting it from its zip into work if needed."""
    if os.path.exists(stencil_file):
      return stencil_file
    # Zipped, usually <file>.zip, sometimes named without the stencil extension
    zip_file = next(
        (z for z in (stencil_file + ".zip", os.path.splitext(stencil_file)[0] + ".zip")
         if os.path.exists(z)), None)
    if zip_file is None:
      sys.exit(f"Stencil file not found: '{stencil_file}' (and no zip to extract it from)\n"
               "See the README for how to obtain vendor stencils")
    self.log(f"Extracting {zip_file}")
    with zipfile.ZipFile(zip_file) as z:
      names = [n for n in z.namelist()
               if not n.startswith("__MACOSX") and n.lower().endswith((".vss", ".vssx", ".vsx"))]
      if not names:
        sys.exit(f"No stencil file found inside '{zip_file}'")
      return z.extract(names[0], work)

  def extract(self, stencil_file, out_dir, stem):
    os.makedirs(out_dir, exist_ok=True)
    work = tempfile.mkdtemp(prefix="devices_")
    try:
      source = self.unpack(stencil_file, work)
      self.log(f"Splitting {source} into SVGs (devon)")
      cmd = [sys.executable, self.find_devon(), "--input", source, "--split",
             "--output", out_dir, "--to", "svg"]
      result = subprocess.run(cmd, capture_output=True, text=True)
      if result.returncode != 0:
        sys.exit(f"devon failed on '{source}':\n{result.stdout}{result.stderr}")
    finally:
      shutil.rmtree(work, ignore_errors=True)

  # ----------------------------------------------------------------------
  # Discovery: find a stencil for a vendor/model from the directory layout
  # (visio-stencils/<first letter>/<vendor>/<stencil>.vss[x].zip)
  # ----------------------------------------------------------------------

  def vendor_dirs(self, vendor):
    want = norm(VENDOR_ALIASES.get(norm(vendor), vendor))
    found = []
    if len(want) < 2 or not os.path.isdir(self.stencil_dir):
      return found
    for letter in sorted(os.listdir(self.stencil_dir)):
      letter_dir = os.path.join(self.stencil_dir, letter)
      if len(letter) != 1 or not os.path.isdir(letter_dir):
        continue
      for name in sorted(os.listdir(letter_dir)):
        n = norm(name)
        if os.path.isdir(os.path.join(letter_dir, name)) and (n == want or (len(n) >= 3 and (n.startswith(want) or want.startswith(n)))):
          found.append(os.path.join(letter, name))
    return found

  def vendor_stencils(self, vendor, model=""):
    """Stencil files (relative paths) for a vendor, most likely first: one per stencil name,
    preferring .vssx (cheap to index), then current over classic, then server-ish names."""
    by_stem = {}
    for vdir in self.vendor_dirs(vendor):
      for name in sorted(os.listdir(os.path.join(self.stencil_dir, vdir))):
        m = re.match(r"(.+?)(\.(?:vss|vssx|vsx))?(?:\.zip)?$", name, re.IGNORECASE)
        if not m or not name.lower().endswith((".zip", ".vss", ".vssx", ".vsx")):
          continue
        stem, ext = m.group(1), (m.group(2) or ".vss")
        rel = f"{vdir}/{stem}{ext}".replace(os.sep, "/")
        old = by_stem.get(stem.lower())
        if old is None or (ext.lower() == ".vssx" and not old.lower().endswith(".vssx")):
          by_stem[stem.lower()] = rel

    series = re.match(r"[a-z]{2,}", norm(model))   # DL380 -> dl, found in HPE-ProLiant-DL.vssx

    def rank(rel):
      base = os.path.basename(rel).lower()
      return (
          not (series and re.search(rf"(?<![a-z]){series.group(0)}(?![a-z])", base)),
          bool(re.search(r"icon|3d|logical|clipart|template|cable|desktop|monitor|cluster|network|software|logo", base)),
          "classic" in base,
          not re.search(r"server|rack|blade|compute|storage|disk|array|series|node|system", base),
          not base.endswith(".vssx"),
          base,
      )
    return sorted(by_stem.values(), key=rank)

  def master_names(self, rel):
    """Master names in a stencil, cached in <cachedir>/_index so each stencil is only read once."""
    index = os.path.join(self.cache_dir, "_index", re.sub(r"[^A-Za-z0-9._-]+", "_", rel) + ".json")
    if os.path.exists(index):
      with open(index) as f:
        return json.load(f)
    work = tempfile.mkdtemp(prefix="devices_idx_")
    try:
      source = self.unpack(os.path.join(self.stencil_dir, *rel.split("/")), work)
      if source.lower().endswith(".vssx"):
        with zipfile.ZipFile(source) as z:
          xml = z.read("visio/masters/masters.xml").decode("utf-8", errors="replace")
        names = re.findall(r"<Master\b[^>]*?\bNameU='([^']*)'", xml) or re.findall(r'<Master\b[^>]*?\bNameU="([^"]*)"', xml)
        names = [n.replace("&amp;", "&").replace("&apos;", "'").replace("&quot;", '"') for n in names]
      else:
        tool = shutil.which("vss2raw")
        if tool is None:
          return []
        out = subprocess.run([tool, source], capture_output=True).stdout.decode("utf-8", errors="replace")
        names = re.findall(r"startPage\(draw:name: (.*?), svg:height: [\d.]+in, svg:width: [\d.]+in\)", out)
    finally:
      shutil.rmtree(work, ignore_errors=True)
    os.makedirs(os.path.dirname(index), exist_ok=True)
    with open(index, "w") as f:
      json.dump(names, f)
    return names

  def discover(self, vendor, model):
    """Search the vendor's stencils for front/rear masters for a model.
    Returns (stencil key, front master, rear master or None), or None."""
    memo = (norm(vendor), norm(model))
    if memo in self.found:
      return self.found[memo]
    result = None
    if len(norm(model)) >= 2:
      candidates = self.vendor_stencils(vendor, model)[:self.max_scan]
      if candidates:
        print(f"Searching {len(candidates)} {vendor} stencil(s) for model '{model}'...", file=sys.stderr)
      for rel in candidates:
        try:
          names = self.master_names(rel)
        except (zipfile.BadZipFile, KeyError, OSError):
          continue
        front = best_master(names, model, ("front",))
        if front:
          rear = best_master(names, model, ("rear", "back"))
          self.log(f"Found '{front}' in {rel}")
          result = (rel, front, rear)
          break
      if result is None:
        reason = f"no stencil found for {vendor} '{model}'" if candidates else f"no '{vendor}' directory in {self.stencil_dir}"
        print(f"Warning: {reason}", file=sys.stderr)
    self.found[memo] = result
    return result

  def find_svg(self, key, master):
    directory = self.svg_dir(key)
    wanted = shape_filename(master)
    path = os.path.join(directory, wanted)
    if os.path.exists(path):
      return path
    lower = wanted.lower()
    for name in os.listdir(directory):
      if name.lower() == lower:
        return os.path.join(directory, name)
    return None

  def load(self, key, master):
    """Return (svg root, width in, height in) for a master, or None if it has no SVG."""
    path = self.find_svg(key, master)
    if path is None:
      return None
    if path not in self.parsed:
      root = ET.parse(path).getroot()
      self.parsed[path] = (root,) + svg_size(root)
    return self.parsed[path]


def to_inches(value):
  m = re.match(r"\s*([0-9.]+)\s*(in|pt|px|cm|mm)?", value or "")
  if not m:
    return None
  num = float(m.group(1))
  return num / {"in": 1, "pt": 72, "px": 96, "cm": 2.54, "mm": 25.4, None: 96}[m.group(2)]


def svg_size(root):
  w = to_inches(root.get("width"))
  h = to_inches(root.get("height"))
  if w is None or h is None:
    vb = [float(v) for v in root.get("viewBox").replace(",", " ").split()]
    w, h = vb[2] / 72.0, vb[3] / 72.0
  return w, h


# --------------------------------------------------------------------------
# SVG composition
# --------------------------------------------------------------------------

_instance = [0]


def embed(parent, root, x_in, y_in, w_in, h_in):
  """Nest a copy of a shape SVG into parent at x,y (inches, top-left origin)."""
  _instance[0] += 1
  prefix = f"i{_instance[0]}_"
  node = copy.deepcopy(root)
  # Shapes reuse ids (Layer1271, gradients, clip paths); make them unique per instance
  ids = {e.get("id") for e in node.iter() if e.get("id")}
  if ids:
    pattern = re.compile(r"(url\(#|#)(" + "|".join(re.escape(i) for i in sorted(ids, key=len, reverse=True)) + r")\b")
    for e in node.iter():
      for attr, val in list(e.attrib.items()):
        if attr == "id":
          e.set(attr, prefix + val)
        elif "#" in val:
          e.set(attr, pattern.sub(lambda m: m.group(1) + prefix + m.group(2), val))
  if not node.get("viewBox"):
    node.set("viewBox", f"0 0 {w_in * 72:.4f} {h_in * 72:.4f}")
  node.set("x", f"{x_in * UNITS:.3f}")
  node.set("y", f"{y_in * UNITS:.3f}")
  node.set("width", f"{w_in * UNITS:.3f}")
  node.set("height", f"{h_in * UNITS:.3f}")
  node.set("preserveAspectRatio", "none")
  node.set("overflow", "hidden")
  parent.append(node)
  return node


def add_label(parent, text, x_in, y_in, rotate=False):
  """Orange tag with the label text. x,y is the bottom-left corner (inches);
  when rotated it reads bottom to top, with x,y the bottom-left of the rotated box."""
  size = LABEL_FONT_PT / 72.0 * UNITS
  width = len(text) * size * 0.52 + size * 0.6
  height = size * 1.45
  x, y = x_in * UNITS, y_in * UNITS
  g = ET.SubElement(parent, f"{{{SVG_NS}}}g")
  if rotate:
    g.set("transform", f"translate({x:.3f},{y:.3f}) rotate(-90)")
  else:
    g.set("transform", f"translate({x:.3f},{y - height:.3f})")
  rect = ET.SubElement(g, f"{{{SVG_NS}}}rect", width=f"{width:.3f}", height=f"{height:.3f}", fill=LABEL_FILL)
  del rect
  t = ET.SubElement(g, f"{{{SVG_NS}}}text", x=f"{size * 0.3:.3f}", y=f"{height * 0.72:.3f}")
  t.set("font-family", "Arial, Helvetica, sans-serif")
  t.set("font-size", f"{size:.3f}")
  t.set("fill", LABEL_TEXT)
  t.text = text
  return width / UNITS


def build_rack_page(lib, rack_name, rows, opts):
  """Compose one page (front and rear elevation of a rack) as an SVG element tree."""
  rack = lib.load("dell_rack", DEFAULT_RACK)
  if rack is None:
    sys.exit(f"Rack frame '{DEFAULT_RACK}' has no SVG")
  rack_root, rack_w, rack_h = rack

  top = MARGIN_Y + (PAGE_LABEL_HEIGHT if opts.pagelabels else 0)
  page_w = MARGIN_X + RACK_PITCH + rack_w + MARGIN_X
  page_h = top + rack_h + MARGIN_Y

  svg = ET.Element(f"{{{SVG_NS}}}svg", {
      "width": f"{page_w:.4f}in", "height": f"{page_h:.4f}in",
      "viewBox": f"0 0 {page_w * UNITS:.3f} {page_h * UNITS:.3f}",
  })
  ET.SubElement(svg, f"{{{SVG_NS}}}rect", width="100%", height="100%", fill="#ffffff")

  if opts.pagelabels:
    t = ET.SubElement(svg, f"{{{SVG_NS}}}text", x=f"{page_w * UNITS / 2:.3f}", y=f"{(MARGIN_Y + 0.35) * UNITS:.3f}")
    t.set("text-anchor", "middle")
    t.set("font-family", "Arial, Helvetica, sans-serif")
    t.set("font-size", f"{16 / 72 * UNITS:.3f}")
    t.set("font-weight", "bold")
    t.text = rack_name

  placed = []
  for row in rows:
    key, front_name, back_name = pick_shape(row)
    names = [front_name, back_name]
    shapes = [lib.load(key, m) for m in names]
    if (None in shapes or front_name == BLANK_PLATE) and opts.discover:
      # No rule for this vendor/model, or its master isn't in the rule's stencil: search the vendor's stencils
      found = lib.discover(row["Vendor"], row["Model"])
      if found:
        key, front_name, back_name = found
        names = [front_name, back_name]
        shapes = [lib.load(key, front_name), lib.load(key, back_name) if back_name else None]
    for i, shape in enumerate(shapes):
      if shape is None:
        print(f"Warning: no SVG for '{names[i]}' ({row['Hostname']}: {row['Component']}); using blank plate",
              file=sys.stderr)
    blank = lib.load("dell_rack", BLANK_PLATE)
    if shapes[0] is None:
      shapes[0] = blank
    if shapes[1] is None:
      # Missing view: a blank plate stretched to the other view's size
      shapes[1] = (blank[0], shapes[0][1], shapes[0][2])
    placed.append((row, shapes))

  for side, offset in (("front", 0), ("rear", 1)):
    rack_x = MARGIN_X + RACK_PITCH * offset
    embed(svg, rack_root, rack_x, top, rack_w, rack_h)
    rack_bottom = top + rack_h
    labels = []
    for row, shapes in placed:
      root, w, h = shapes[offset]
      rus = to_float(row["Rack Units"])
      top_ru = to_float(row["Top Rack Unit"])
      x = rack_x + (rack_w - w) / 2
      bottom = rack_bottom - RU_BASE - (top_ru - rus) * RU_SPACE
      y = bottom - h
      embed(svg, root, x, y, w, h)
      labels.append((f"{row['Hostname']}: {row['Component']}", x, bottom))
    if opts.showlabels:
      for text, x, bottom in labels:
        add_label(svg, text, x + 0.03, bottom - 0.02)
      label_w = len(rack_name) * LABEL_FONT_PT / 72.0 * 0.52 + LABEL_FONT_PT / 72.0 * 0.6
      add_label(svg, rack_name, rack_x - LABEL_FONT_PT / 72.0 * 1.45, top + (rack_h + label_w) / 2, rotate=True)
  return svg


COLUMNS = ["Hostname", "Component", "Vendor", "Architecture", "Model", "Operating System", "Rack",
           "Rack Units", "Top Rack Unit", "Serial Number", "Asset Number", "Installed Date",
           "Warranty Exp", "Location", "Country"]


def cell_text(value):
  """Text for a spreadsheet cell: whole numbers without the .0, dates as YYYY-MM-DD."""
  if value is None:
    return ""
  if isinstance(value, bool):
    return str(value)
  if isinstance(value, float) and value.is_integer():
    return str(int(value))
  if isinstance(value, datetime.datetime):
    return value.date().isoformat() if value.time() == datetime.time() else value.isoformat(sep=" ")
  if isinstance(value, datetime.date):
    return value.isoformat()
  return str(value).strip()


def read_csv(path):
  with open(path, newline="", encoding="utf-8-sig") as f:
    return [[c for c in row] for row in csv.reader(f)]


def read_xlsx(path, sheet):
  try:
    import openpyxl
  except ImportError:
    sys.exit("openpyxl is needed to read .xlsx files (pip install openpyxl)")
  book = openpyxl.load_workbook(path, read_only=True, data_only=True)
  try:
    ws = book[sheet] if sheet else book.worksheets[0]
  except KeyError:
    sys.exit(f"Sheet '{sheet}' not found in {path}; sheets: {', '.join(book.sheetnames)}")
  return [[cell_text(v) for v in row] for row in ws.iter_rows(values_only=True)]


def read_xls(path, sheet):
  try:
    import xlrd
  except ImportError:
    sys.exit("xlrd is needed to read .xls files (pip install xlrd)")
  book = xlrd.open_workbook(path)
  try:
    ws = book.sheet_by_name(sheet) if sheet else book.sheet_by_index(0)
  except xlrd.XLRDError:
    sys.exit(f"Sheet '{sheet}' not found in {path}; sheets: {', '.join(book.sheet_names())}")
  rows = []
  for r in range(ws.nrows):
    row = []
    for cell in ws.row(r):
      if cell.ctype == xlrd.XL_CELL_DATE:
        row.append(cell_text(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode)))
      elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
        row.append("")
      else:
        row.append(cell_text(cell.value))
    rows.append(row)
  return rows


def load_rows(path, sheet=None):
  """Read hardware rows from a CSV, xlsx or xls file (first sheet unless sheet is given).
  The first row is the header; columns are matched by name, ignoring case and spacing."""
  ext = os.path.splitext(path)[1].lower()
  if ext not in (".csv", ".xlsx", ".xlsm", ".xls"):
    # Unknown extension: go by the file's first bytes (zip = xlsx, OLE2 = xls), else treat as CSV
    with open(path, "rb") as f:
      magic = f.read(4)
    ext = ".xlsx" if magic[:2] == b"PK" else ".xls" if magic == b"\xd0\xcf\x11\xe0" else ".csv"
  if ext == ".csv":
    table = read_csv(path)
  elif ext == ".xls":
    table = read_xls(path, sheet)
  else:
    table = read_xlsx(path, sheet)
  if not table:
    sys.exit(f"No data in '{path}'")
  header = {norm(h): i for i, h in enumerate(table[0])}
  missing = [c for c in ("Hostname", "Rack", "Rack Units", "Top Rack Unit") if norm(c) not in header]
  if missing:
    sys.exit(f"'{path}' has no {', '.join(missing)} column(s); the first row must be the column headings")
  rows = []
  for values in table[1:]:
    if not any(str(v).strip() for v in values):
      continue
    # Short rows leave later columns missing; treat them as empty
    rows.append({c: (str(values[header[norm(c)]]).strip() if norm(c) in header and header[norm(c)] < len(values) else "")
                 for c in COLUMNS})
  return rows


def to_float(value):
  try:
    return float(value)
  except (TypeError, ValueError):
    return 0.0


# --------------------------------------------------------------------------
# Output
# --------------------------------------------------------------------------

def write_svg(svg, path):
  ET.ElementTree(svg).write(path, encoding="utf-8", xml_declaration=True)


def render(svg_files, output, fmt, dpi):
  """Rasterise/convert SVG page files with rsvg-convert. pdf takes all pages."""
  rsvg = shutil.which("rsvg-convert")
  if rsvg is None:
    sys.exit("rsvg-convert not found (install librsvg) to write png/jpg/pdf output")
  if fmt == "pdf":
    subprocess.run([rsvg, "-f", "pdf", "-o", output] + svg_files, check=True)
    return
  png = output if fmt == "png" else output + ".png"
  subprocess.run([rsvg, "-f", "png", "-d", str(dpi), "-p", str(dpi), "-b", "white", "-o", png, svg_files[0]], check=True)
  if fmt == "jpg":
    try:
      from PIL import Image
    except ImportError:
      sys.exit("Pillow is needed for jpg output (pip install Pillow)")
    Image.open(png).convert("RGB").save(output, quality=92)
    os.remove(png)


def safe_name(name):
  return re.sub(r'[\\/:*?"<>|]', "_", name)


def main():
  p = argparse.ArgumentParser(
      description="Generate rack elevation diagrams from a CSV file, using SVGs extracted from Visio stencils",
      allow_abbrev=False)
  p.add_argument("-inputfile", "--inputfile", metavar="FILENAME", help="CSV, xls or xlsx file describing the hardware")
  p.add_argument("-outputfile", "--outputfile", metavar="FILENAME",
                 help="output file; format from extension: .svg .png .jpg .pdf")
  p.add_argument("-sheet", "--sheet", metavar="NAME", help="worksheet to read from an .xls/.xlsx file (default: the first)")
  p.add_argument("-longracknames", "--longracknames", action="store_true",
                 help="append chassis hostnames to rack names")
  p.add_argument("-showlabels", "--showlabels", action="store_true", help="show hostname labels")
  p.add_argument("-rackperfile", "--rackperfile", action="store_true",
                 help="write one file per rack into the output directory")
  p.add_argument("-pagelabels", "--pagelabels", action="store_true", help="draw the rack name at the top of each page")
  p.add_argument("-stencildir", "--stencildir", default=os.path.join(SCRIPT_DIR, "visio-stencils"),
                 help="visio-stencils checkout (default: visio-stencils next to this script)")
  p.add_argument("-cachedir", "--cachedir", default=os.path.join(SCRIPT_DIR, "svg-cache"),
                 help="where extracted stencil SVGs are cached (default: svg-cache next to this script)")
  p.add_argument("-devon", "--devon", help="path to devon.py (default: $DEVON or ../devon/devon.py)")
  p.add_argument("-nodiscover", "--nodiscover", dest="discover", action="store_false",
                 help="don't search the visio-stencils directory for models with no built-in rule")
  p.add_argument("-maxscan", "--maxscan", type=int, default=30,
                 help="most stencils to search per vendor/model when discovering (default 30)")
  p.add_argument("-dpi", "--dpi", type=int, default=150, help="resolution for png/jpg output (default 150)")
  p.add_argument("-verbose", "--verbose", action="store_true")
  p.add_argument("-version", "--version", action="store_true")
  opts = p.parse_args()

  if opts.version:
    print(get_script_vers())
    return
  if not opts.inputfile:
    p.error("input file not specified")
  if not os.path.exists(opts.inputfile):
    sys.exit(f"File: '{opts.inputfile}' does not exist")
  if not opts.outputfile and not opts.rackperfile:
    p.error("output file not specified")

  rows = load_rows(opts.inputfile, opts.sheet)

  racks = list(dict.fromkeys(r["Rack"] for r in rows))
  lib = StencilLibrary(opts.stencildir, opts.cachedir, opts.devon, opts.verbose, opts.maxscan)

  pages = []
  for rack in racks:
    rack_rows = [r for r in rows if r["Rack"] == rack]
    rack_name = rack
    if opts.longracknames:
      hosts = ",".join(r["Hostname"] for r in rack_rows if csi("CH|Chassis", r["Component"]))
      rack_name = f"{rack} ({hosts})"
    pages.append((rack_name, build_rack_page(lib, rack_name, rack_rows, opts)))

  out_dir = os.path.join(SCRIPT_DIR, "output")
  work = tempfile.mkdtemp(prefix="devices_pages_")
  try:
    if opts.rackperfile:
      os.makedirs(out_dir, exist_ok=True)
      fmt = os.path.splitext(opts.outputfile)[1].lstrip(".").lower() if opts.outputfile else "svg"
      for rack_name, svg in pages:
        path = os.path.join(out_dir, f"{safe_name(rack_name)}.{fmt or 'svg'}")
        emit([svg], path, fmt or "svg", opts.dpi, work)
        print(f"Wrote {path}")
    else:
      output = opts.outputfile
      fmt = os.path.splitext(output)[1].lstrip(".").lower()
      if fmt not in ("svg", "png", "jpg", "jpeg", "pdf"):
        sys.exit(f"Unsupported output format '.{fmt}': use .svg, .png, .jpg or .pdf")
      fmt = "jpg" if fmt == "jpeg" else fmt
      if len(pages) == 1 or fmt == "pdf":
        emit([s for _, s in pages], output, fmt, opts.dpi, work)
        print(f"Wrote {output}")
      else:
        # Single-page formats get one file per rack
        base, ext = os.path.splitext(output)
        for rack_name, svg in pages:
          path = f"{base}_{safe_name(rack_name)}{ext}"
          emit([svg], path, fmt, opts.dpi, work)
          print(f"Wrote {path}")
  finally:
    shutil.rmtree(work, ignore_errors=True)


def emit(svgs, path, fmt, dpi, work):
  os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
  if fmt == "svg":
    write_svg(svgs[0], path)
    return
  files = []
  for i, svg in enumerate(svgs):
    f = os.path.join(work, f"page{i}.svg")
    write_svg(svg, f)
    files.append(f)
  render(files, path, fmt, dpi)


if __name__ == "__main__":
  main()
